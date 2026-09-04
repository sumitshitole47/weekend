import os
import re
import shutil
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, fill_hex: str):
    """Set background color of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal padding for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_formatted_runs(paragraph, text: str):
    """Parse inline **bold** and `code` formatting and append runs to paragraph."""
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|[^\*`]+)')
    tokens = pattern.findall(text)

    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            paragraph.add_run(token)


def markdown_to_docx(md_path: str, docx_path: str):
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"Markdown file not found: {md_path}")

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Set document margins (0.8 inches)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Set default Normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(36, 41, 46)

    lines = md_text.split("\n")
    in_table = False
    table_headers = []
    table_rows = []

    for line in lines:
        line_str = line.strip()

        # Heading 1
        if line_str.startswith("# "):
            h = doc.add_heading(level=1)
            run = h.add_run(line_str[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.bold = True
            run.font.color.rgb = RGBColor(3, 102, 214)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            continue

        # Heading 2
        if line_str.startswith("## "):
            h = doc.add_heading(level=2)
            run = h.add_run(line_str[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(47, 54, 61)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            continue

        # Heading 3
        if line_str.startswith("### "):
            h = doc.add_heading(level=3)
            run = h.add_run(line_str[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(36, 41, 46)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            continue

        # Blockquote / Callout Box
        if line_str.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, line_str[2:])
            for r in p.runs:
                r.font.italic = True
                r.font.color.rgb = RGBColor(40, 40, 40)
            continue

        # Tables
        if "|" in line_str and line_str.startswith("|"):
            if "---" in line_str:
                continue
            cells = [c.strip() for c in line_str.split("|")[1:-1]]
            if not in_table:
                in_table = True
                table_headers = cells
            else:
                table_rows.append(cells)
            continue
        else:
            if in_table:
                # Render accumulated table
                table = doc.add_table(rows=1 + len(table_rows), cols=len(table_headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Format Header
                hdr_cells = table.rows[0].cells
                for i, th in enumerate(table_headers):
                    hdr_cells[i].text = th
                    set_cell_background(hdr_cells[i], "F6F8FA")
                    for p in hdr_cells[i].paragraphs:
                        for r in p.runs:
                            r.bold = True
                            r.font.color.rgb = RGBColor(3, 102, 214)

                # Format Data Rows
                for r_idx, r_data in enumerate(table_rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8F9FA"
                    for c_idx, val in enumerate(r_data):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = val
                            set_cell_background(row_cells[c_idx], bg_color)

                doc.add_paragraph()  # spacing after table
                in_table = False
                table_headers = []
                table_rows = []

        # Bullet List
        if line_str.startswith("- ") or line_str.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, line_str[2:])
            continue

        # Code block marker
        if line_str.startswith("```"):
            continue

        # Normal Paragraph
        if line_str:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, line_str)

    if in_table:
        table = doc.add_table(rows=1 + len(table_rows), cols=len(table_headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        hdr_cells = table.rows[0].cells
        for i, th in enumerate(table_headers):
            hdr_cells[i].text = th
            set_cell_background(hdr_cells[i], "F6F8FA")
            for p in hdr_cells[i].paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(3, 102, 214)

        for r_idx, r_data in enumerate(table_rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8F9FA"
            for c_idx, val in enumerate(r_data):
                if c_idx < len(row_cells):
                    row_cells[c_idx].text = val
                    set_cell_background(row_cells[c_idx], bg_color)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"[SUCCESS] DOCX generated successfully: '{docx_path}'")


def main():
    md_path = r"C:\Users\shito\.gemini\antigravity\brain\1f6c915a-0caa-46be-840a-738018d735f1\presentation_guide.md"
    docx_artifact_path = r"C:\Users\shito\.gemini\antigravity\brain\1f6c915a-0caa-46be-840a-738018d735f1\presentation_guide.docx"
    docx_project_path = r"C:\Users\shito\.gemini\antigravity\scratch\sih26158_drone_3d\presentation_guide.docx"

    markdown_to_docx(md_path, docx_project_path)
    shutil.copy(docx_project_path, docx_artifact_path)
    print(f"[SAVED] Artifact copy saved to: '{docx_artifact_path}'")


if __name__ == "__main__":
    main()
